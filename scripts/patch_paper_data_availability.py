#!/usr/bin/env python3
"""Update Data Availability in papers.docx: supplementary ZIP, not GitHub."""

from datetime import datetime
from pathlib import Path
import shutil

from docx import Document

SRC = Path(r"E:\Chj_projects\Papers\papers.docx")
BACKUP = Path(r"E:\Chj_projects\Papers") / (
    f"papers_backup_avail_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
)

NEW = (
    "Data Availability Statement: The source code, the Monte Carlo "
    "initial-condition generator, the protocol files, and the simulated "
    "engagement .mat files used in this study are provided as Supplementary "
    "Material (ZIP archive) with this article. The engagement conditions are "
    "sampled from published Sparrow-class (AIM-7F) parameters with fixed "
    "random seeds; they are not recorded flight data. Bootstrap-resampled or "
    "post-hoc synthetic result rows used only for internal diagnostics are "
    "excluded from the archive and were not used in the reported tables. "
        "Exact API model identifiers and experimental dates (GSM campaign "
        "2-12 August 2026; equal-budget baselines 20-22 August 2026, UTC+8) "
        "are listed in REPRODUCIBILITY.md inside the archive."
)


def set_text(para, new_text: str) -> None:
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(new_text)


def main() -> None:
    shutil.copy2(SRC, BACKUP)
    doc = Document(str(SRC))
    n = 0
    for para in doc.paragraphs:
        t = para.text
        if "Data Availability" in t or "github.com/u3693025048-ops/GSMultiAgent2" in t:
            set_text(para, NEW)
            n += 1
            break
    if n == 0:
        raise SystemExit("Data Availability paragraph not found")
    doc.save(str(SRC))
    print(f"updated {n} paragraph")
    print(f"backup {BACKUP}")


if __name__ == "__main__":
    main()

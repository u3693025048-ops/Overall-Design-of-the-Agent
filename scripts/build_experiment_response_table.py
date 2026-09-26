#!/usr/bin/env python3
"""Build 1:1 response table for 实验清单统计表.xlsx from local (+ Ubuntu) deliverables."""

from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "results" / "experiment_plan"
SRC = Path(r"c:\Users\admin\Desktop\实验清单统计表.xlsx")
OUT_DESKTOP = Path(r"c:\Users\admin\Desktop\实验清单响应表.xlsx")
OUT_REPO = PLAN / "实验清单响应表.xlsx"


def read_csv(path: Path) -> List[Dict[str, str]]:
    if not path.is_file() or path.stat().st_size == 0:
        return []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def gsm_info(paths: Sequence[str], *, methods: Optional[Sequence[str]] = None) -> Dict[str, Any]:
    rows: List[Dict[str, str]] = []
    used: List[str] = []
    for rel in paths:
        p = ROOT / rel if not Path(rel).is_absolute() else Path(rel)
        part = read_csv(p)
        if part:
            used.append(str(p.relative_to(ROOT)).replace("\\", "/") if p.is_relative_to(ROOT) else str(p))
            rows.extend(part)
    if methods:
        allow = set(methods)
        rows = [r for r in rows if r.get("method") in allow]
    if not rows:
        return {"n": 0, "paths": [], "seeds": [], "scens": [], "snap": 0, "empty_dev": 0, "note": ""}
    seeds = sorted({str(r.get("seed")) for r in rows})
    scens = sorted({str(r.get("scenario")) for r in rows})
    meths = sorted({str(r.get("method")) for r in rows})
    snap = sum(1 for r in rows if "without fresh metrics snapshot" in (r.get("error_msg") or ""))
    empty_dev = sum(
        1
        for r in rows
        if str(r.get("SEP_dev", "")).lower() in ("inf", "infinity")
        and str(r.get("hit_rate_dev")) in ("0", "0.0", "")
    )
    return {
        "n": len(rows),
        "paths": used,
        "seeds": seeds,
        "scens": scens,
        "methods": meths,
        "snap": snap,
        "empty_dev": empty_dev,
        "note": f"{len(rows)}行 | methods={','.join(meths)} | seeds={','.join(seeds)} | scenarios={','.join(scens)} | snap缺失={snap} | empty_dev={empty_dev}",
    }


def opt_info(dirs: Sequence[str], method: str) -> Dict[str, Any]:
    rows: List[Dict[str, str]] = []
    used: List[str] = []
    json_n = 0
    for rel in dirs:
        d = ROOT / rel
        if not d.is_dir():
            continue
        used.append(rel)
        rows.extend(read_csv(d / "runs.csv"))
        if (d / "runs").is_dir():
            json_n += len(list((d / "runs").glob("*.json")))
    rows = [r for r in rows if (r.get("method") or method) == method or method in (r.get("method") or "")]
    # some CSVs only have one method per dir
    if not rows:
        # retry without filter if dir is method-specific
        rows = []
        for rel in dirs:
            d = ROOT / rel
            part = read_csv(d / "runs.csv")
            rows.extend(part)
    seeds = sorted(
        {
            str(r.get("master_seed") or r.get("seed") or "")
            for r in rows
            if (r.get("master_seed") or r.get("seed"))
        }
    )
    # parse from json names if seed missing
    if not seeds:
        for rel in dirs:
            rd = ROOT / rel / "runs"
            if rd.is_dir():
                for p in rd.glob("*.json"):
                    if "_seed" in p.stem:
                        seeds.append(p.stem.split("_seed")[-1])
        seeds = sorted(set(seeds))
    scens = sorted({str(r.get("scenario")) for r in rows if r.get("scenario")})
    if not scens:
        for rel in dirs:
            rd = ROOT / rel / "runs"
            if rd.is_dir():
                for p in rd.glob("*.json"):
                    scens.append(p.stem.split("_")[0])
        scens = sorted(set(scens))
    n = max(len(rows), json_n)
    return {
        "n": n,
        "paths": used,
        "seeds": seeds,
        "scens": scens,
        "json_n": json_n,
        "note": f"csv/json约{n} | seeds={','.join(seeds) or '?'} | scenarios={','.join(scens)} | dirs={'; '.join(used)}",
    }


def pull_ubuntu_extra() -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    try:
        import paramiko
    except Exception:
        return out
    pass_path = ROOT / "scripts" / ".gsm_ssh_pass"
    if not pass_path.is_file():
        return out
    password = pass_path.read_text(encoding="utf-8").strip()
    c = paramiko.SSHClient()
    c.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        c.connect(
            "10.13.11.221",
            username="admcai",
            password=password,
            timeout=25,
            allow_agent=False,
            look_for_keys=False,
        )
    except Exception as exc:
        out["error"] = str(exc)
        return out
    cmd = r"""
python3 - <<'PY'
import csv, json
from pathlib import Path
plan=Path('/home/admcai/workspace/GSMultiAgent13-V0.3/results/experiment_plan')
for name in ['ablation_correction_ubuntu_42_seed42.csv','ablation_optimization_ubuntu_42_seed42.csv']:
 p=plan/name
 if not p.exists() or p.stat().st_size==0:
  print(name, 'EMPTY')
  continue
 rows=list(csv.DictReader(p.open(encoding='utf-8-sig')))
 print(name, 'N', len(rows))
 for r in rows:
  print(name, r.get('method'), r.get('scenario'), r.get('SEP_dev'), r.get('error_msg','')[:40])
st=json.loads((plan/'ubuntu_queue_status.json').read_text())
print('QUEUE', st.get('phase'), st.get('current_job'), st.get('last_message'))
PY
"""
    _, o, _ = c.exec_command(cmd, timeout=60)
    text = o.read().decode("utf-8", "replace")
    c.close()
    out["raw"] = text
    # parse light
    corr_n = 0
    opt_n = 0
    for line in text.splitlines():
        if line.startswith("ablation_correction_ubuntu_42_seed42.csv N"):
            corr_n = int(line.split()[-1])
        if line.startswith("ablation_optimization_ubuntu_42_seed42.csv N"):
            opt_n = int(line.split()[-1])
    out["corr42_n"] = corr_n
    out["opt42_n"] = opt_n
    return out


def status_for(need_seeds: Sequence[str], have_seeds: Sequence[str], n: int, expect_n: int, *, issues: str = "") -> str:
    if n <= 0:
        return "缺失"
    have = set(have_seeds)
    need = set(need_seeds)
    if issues:
        return "部分完整"
    if need and not need.issubset(have):
        return "部分完整"
    if expect_n and n < expect_n:
        return "部分完整"
    return "完整"


def main() -> int:
    ubu = pull_ubuntu_extra()

    # ---- gather ----
    e0 = gsm_info(
        [
            "results/experiment_plan/ablation_memory.csv",
            "results/experiment_plan/ablation_memory_seed42.csv",
            "results/experiment_plan/ablation_memory_ubuntu_43_44.csv",
            "results/experiment_plan/ablation_memory_ubuntu.csv",
        ],
        methods=["E0_full"],
    )
    # dedupe e0 by seed+scenario from preferred files
    e0_rows = []
    for rel in [
        "results/experiment_plan/ablation_memory_seed42.csv",
        "results/experiment_plan/ablation_memory.csv",
        "results/experiment_plan/ablation_memory_ubuntu_43_44.csv",
    ]:
        for r in read_csv(ROOT / rel):
            if r.get("method") == "E0_full":
                e0_rows.append(r)
    # unique by seed,scenario keep first
    seen = set()
    e0_uniq = []
    for r in e0_rows:
        k = (str(r.get("seed")), str(r.get("scenario")))
        if k in seen:
            continue
        seen.add(k)
        e0_uniq.append(r)
    e0 = {
        "n": len(e0_uniq),
        "paths": [
            "results/experiment_plan/ablation_memory_seed42.csv",
            "results/experiment_plan/ablation_memory_ubuntu_43_44.csv",
        ],
        "seeds": sorted({str(r.get("seed")) for r in e0_uniq}),
        "scens": sorted({str(r.get("scenario")) for r in e0_uniq}),
        "snap": sum(1 for r in e0_uniq if "without fresh metrics snapshot" in (r.get("error_msg") or "")),
        "empty_dev": sum(
            1
            for r in e0_uniq
            if str(r.get("SEP_dev", "")).lower() in ("inf", "infinity")
            and str(r.get("hit_rate_dev")) in ("0", "0.0", "")
        ),
    }
    e0["note"] = (
        f"{e0['n']}行 | seeds={','.join(e0['seeds'])} | scenarios={','.join(e0['scens'])} | "
        f"snap缺失={e0['snap']} | empty_dev={e0['empty_dev']}"
    )

    e123_42 = gsm_info(
        ["results/experiment_plan/ablation_memory_E123_local_42.csv"],
        methods=["E1_no_memory", "E2_no_reflection", "E3_no_memory_no_reflection"],
    )
    e123_4344 = gsm_info(
        ["results/experiment_plan/ablation_memory_E123_ubuntu_43_44.csv"],
        methods=["E1_no_memory", "E2_no_reflection", "E3_no_memory_no_reflection"],
    )

    def e123_method(method: str) -> Dict[str, Any]:
        rows = []
        paths = []
        for rel in [
            "results/experiment_plan/ablation_memory_E123_local_42.csv",
            "results/experiment_plan/ablation_memory_E123_ubuntu_43_44.csv",
        ]:
            part = [r for r in read_csv(ROOT / rel) if r.get("method") == method]
            if part:
                paths.append(rel)
                rows.extend(part)
        seeds = sorted({str(r.get("seed")) for r in rows})
        snap = sum(1 for r in rows if "without fresh metrics snapshot" in (r.get("error_msg") or ""))
        empty_dev = sum(
            1
            for r in rows
            if str(r.get("SEP_dev", "")).lower() in ("inf", "infinity")
            and str(r.get("hit_rate_dev")) in ("0", "0.0", "")
        )
        return {
            "n": len(rows),
            "paths": paths,
            "seeds": seeds,
            "scens": sorted({str(r.get("scenario")) for r in rows}),
            "snap": snap,
            "empty_dev": empty_dev,
            "note": f"{len(rows)}行 | seeds={','.join(seeds)} | snap缺失={snap} | empty_dev={empty_dev}",
        }

    e1 = e123_method("E1_no_memory")
    e2 = e123_method("E2_no_reflection")
    e3 = e123_method("E3_no_memory_no_reflection")

    order_fwd = e0  # R9 uses E0 forward
    order_rev = gsm_info(
        [
            "results/experiment_plan/ablation_order_reverse_local_42.csv",
            "results/experiment_plan/ablation_order_reverse_ubuntu_43_44.csv",
        ],
        methods=["E0_full_reverse"],
    )
    order_shuf = gsm_info(
        [
            "results/experiment_plan/ablation_order_shuffled_local_42.csv",
            "results/experiment_plan/ablation_order_shuffled_ubuntu_43_44.csv",
        ],
        methods=["E0_full_shuffled"],
    )

    o0 = gsm_info(
        [
            "results/experiment_plan/ablation_optimization_ubuntu_43_44.csv",
            "results/experiment_plan/ablation_optimization_local_42.csv",
        ],
        methods=["O0_full"],
    )
    o1 = gsm_info(
        [
            "results/experiment_plan/ablation_optimization_ubuntu_43_44.csv",
            "results/experiment_plan/ablation_optimization_local_42.csv",
        ],
        methods=["O1_no_optimization_agent"],
    )
    # ubuntu opt42 in progress
    opt42_n = int(ubu.get("opt42_n") or 0)
    if opt42_n:
        o0["note"] += f"；Ubuntu seed42 进行中已落盘 {opt42_n}/8 格（逐场景）"
        o1["note"] += f"；Ubuntu seed42 进行中已落盘 {opt42_n}/8 格（含 O0/O1）"

    corr = gsm_info(
        [
            "results/experiment_plan/ablation_correction_ubuntu_43_44.csv",
            "results/experiment_plan/ablation_correction_local_42.csv",
        ]
    )
    corr42_n = int(ubu.get("corr42_n") or 0)
    if corr42_n:
        corr["note"] += f"；Ubuntu seed42 Correction 另有 {corr42_n} 行（C0/C1/C2×T4）"
        if "42" not in corr["seeds"]:
            corr["seeds"] = sorted(set(corr["seeds"]) | {"42"})

    ppo = opt_info(
        [
            "results/experiment_plan/optimizer_baselines_ppo_local_42",
            "results/experiment_plan/optimizer_baselines_ppo_ubuntu_43_44",
        ],
        "direct_ppo",
    )
    cma = opt_info(
        [
            "results/experiment_plan/optimizer_baselines_cma_local_42",
            "results/experiment_plan/optimizer_baselines_cma_ubuntu_43_44",
        ],
        "cma_es",
    )
    rnd = opt_info(
        [
            "results/experiment_plan/optimizer_baselines_random_local_42",
            "results/experiment_plan/optimizer_baselines_random_ubuntu_43_44",
        ],
        "random_search",
    )
    hermes = opt_info(
        ["results/experiment_plan/optimizer_baselines_hermes_ubuntu_42_43_44"],
        "single_agent_hermes",
    )

    # R mapping
    records = [
        {
            "id": "R1",
            "unit": "GSMultiAgent 完整系统（E0 / 混合 PPO）",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": e0,
            "matlab_hint": "DATA/run_archives/experiments/memory_E0_full_seed{42,43,44}_*/",
            "gap": "" if e0["n"] >= 12 and set(["42", "43", "44"]).issubset(e0["seeds"]) else "缺主种子或行数不足",
        },
        {
            "id": "R2",
            "unit": "优化组件替换 +direct_ppo",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": ppo,
            "matlab_hint": "（optimizer baseline，非 GSM 消融归档）",
            "gap": "" if ppo["n"] >= 12 else f"当前约{ppo['n']}/12",
        },
        {
            "id": "R3",
            "unit": "优化组件替换 +cma_es",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": cma,
            "matlab_hint": "（optimizer baseline）",
            "gap": "" if cma["n"] >= 12 else f"当前约{cma['n']}/12",
        },
        {
            "id": "R4",
            "unit": "优化组件替换 +random_search",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": rnd,
            "matlab_hint": "（optimizer baseline）",
            "gap": "" if rnd["n"] >= 12 else f"当前约{rnd['n']}/12",
        },
        {
            "id": "R5",
            "unit": "级联消融 Hermes 单智能体",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": hermes,
            "matlab_hint": "optimizer_baselines_hermes_ubuntu_42_43_44/runs/",
            "gap": "" if hermes["n"] >= 12 else f"当前约{hermes['n']}/12",
        },
        {
            "id": "R6",
            "unit": "记忆消融 E1",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": e1,
            "matlab_hint": "DATA/run_archives/experiments/memory_E1_no_memory_seed*/",
            "gap": ("seed42 有1行 snap/empty_dev" if e1.get("empty_dev") else "")
            or ("" if e1["n"] >= 12 and set(["42", "43", "44"]).issubset(e1["seeds"]) else "行数/种子不全"),
        },
        {
            "id": "R7",
            "unit": "记忆消融 E2",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": e2,
            "matlab_hint": "DATA/run_archives/experiments/memory_E2_no_reflection_seed*/",
            "gap": ("seed42 E2×T4 Layer2 blocked → DEV 快照缺失" if e2.get("empty_dev") else "")
            or ("" if e2["n"] >= 12 else "行数不全"),
        },
        {
            "id": "R8",
            "unit": "记忆消融 E3",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": e3,
            "matlab_hint": "DATA/run_archives/experiments/memory_E3_no_memory_no_reflection_seed*/",
            "gap": "" if e3["n"] >= 12 and set(["42", "43", "44"]).issubset(e3["seeds"]) else "行数/种子不全",
        },
        {
            "id": "R9",
            "unit": "顺序鲁棒性·正向（复用 E0）",
            "need_seeds": ["42"],
            "expect_n": 4,
            "info": order_fwd,
            "matlab_hint": "同 R1 E0 归档",
            "gap": "" if order_fwd["n"] >= 4 else "缺失",
        },
        {
            "id": "R10",
            "unit": "顺序鲁棒性·反向",
            "need_seeds": ["42"],
            "expect_n": 4,
            "info": order_rev,
            "matlab_hint": "DATA/run_archives/experiments/order_E0_full_reverse_seed*/",
            "gap": "" if order_rev["n"] >= 4 else "缺失",
        },
        {
            "id": "R11",
            "unit": "顺序鲁棒性·乱序",
            "need_seeds": ["42"],
            "expect_n": 4,
            "info": order_shuf,
            "matlab_hint": "DATA/run_archives/experiments/order_E0_full_shuffled_seed*/",
            "gap": "" if order_shuf["n"] >= 4 else "缺失",
        },
        {
            "id": "R12",
            "unit": "优化智能体消融 O1（仅 LLM 启发式）",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 12,
            "info": o1,
            "matlab_hint": "DATA/run_archives/experiments/optimization_O1_no_optimization_agent_seed*/",
            "gap": (
                "seed42 本机未落盘；Ubuntu 正在重跑 optimization（含 O0/O1）"
                if "42" not in o1.get("seeds", [])
                else ""
            )
            or ("" if o1["n"] >= 8 else "行数不足"),
        },
        {
            "id": "R13",
            "unit": "修正管线 Layer1 消融 C1",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 3,
            "info": {
                **corr,
                "n": sum(
                    1
                    for rel in [
                        "results/experiment_plan/ablation_correction_ubuntu_43_44.csv",
                    ]
                    for r in read_csv(ROOT / rel)
                    if r.get("method") == "C1_no_layer1_syntax"
                )
                + (1 if corr42_n >= 2 else 0),
            },
            "matlab_hint": "DATA/run_archives/experiments/correction_C1_no_layer1_syntax_seed*/",
            "gap": "",
        },
        {
            "id": "R14",
            "unit": "修正管线 Layers2+3 消融 C2",
            "need_seeds": ["42", "43", "44"],
            "expect_n": 3,
            "info": corr,
            "matlab_hint": "DATA/run_archives/experiments/correction_C2_no_layers23_compat_semantic_seed*/",
            "gap": "",
        },
        {
            "id": "R15",
            "unit": "基础模型替换 Qwen3.7-Max",
            "need_seeds": ["42"],
            "expect_n": 4,
            "info": {"n": 0, "paths": [], "seeds": [], "scens": [], "note": "未交付"},
            "matlab_hint": "—",
            "gap": "未跑 / 无成果物",
        },
        {
            "id": "R16",
            "unit": "基础模型替换 Kimi K3",
            "need_seeds": ["42"],
            "expect_n": 4,
            "info": {"n": 0, "paths": [], "seeds": [], "scens": [], "note": "未交付"},
            "matlab_hint": "—",
            "gap": "未跑 / 无成果物",
        },
    ]

    # refine R13/R14 counts properly
    corr_rows = read_csv(ROOT / "results/experiment_plan/ablation_correction_ubuntu_43_44.csv")
    c0 = [r for r in corr_rows if r.get("method") == "C0_full_pipeline"]
    c1 = [r for r in corr_rows if r.get("method") == "C1_no_layer1_syntax"]
    c2 = [r for r in corr_rows if r.get("method") == "C2_no_layers23_compat_semantic"]
    # R13 is C1, R14 is C2; also C0 is baseline full pipeline for correction chapter
    for rec in records:
        if rec["id"] == "R13":
            seeds = sorted({str(r.get("seed")) for r in c1} | ({"42"} if corr42_n else set()))
            n = len(c1) + (1 if corr42_n else 0)
            rec["info"] = {
                "n": n,
                "paths": ["results/experiment_plan/ablation_correction_ubuntu_43_44.csv"]
                + (["Ubuntu: ablation_correction_ubuntu_42_seed42.csv"] if corr42_n else []),
                "seeds": seeds,
                "scens": ["T4"],
                "note": f"C1 行数={len(c1)}（43/44）+ seed42远程{corr42_n and '有' or '无'}; 协议为T4×主种子，非字面5次独立清空（实现为C0/C1/C2消融）",
            }
            rec["gap"] = "" if n >= 2 else "不足"
            if "42" not in seeds:
                rec["gap"] = (rec["gap"] + "；" if rec["gap"] else "") + "本机无seed42 CSV（远程有则需pull）"
        if rec["id"] == "R14":
            seeds = sorted({str(r.get("seed")) for r in c2} | ({"42"} if corr42_n else set()))
            n = len(c2) + (1 if corr42_n else 0)
            rec["info"] = {
                "n": n,
                "paths": ["results/experiment_plan/ablation_correction_ubuntu_43_44.csv"]
                + (["Ubuntu: ablation_correction_ubuntu_42_seed42.csv"] if corr42_n else []),
                "seeds": seeds,
                "scens": ["T4"],
                "note": f"C2 行数={len(c2)}（43/44）; 另C0基线{len(c0)}行可作对照",
            }
            rec["gap"] = "" if n >= 2 else "不足"

    # O0 note for completeness of R1 vs optimization - R12 uses O1
    # Also add O0 completeness into a side note in R12 gap if needed

    # Build workbook: sheet1 = response 1:1, sheet2 = protocol note
    wb_src = openpyxl.load_workbook(SRC)
    ws_src = wb_src[wb_src.sheetnames[0]]
    # map source rows by id
    src_by_id = {}
    for row in ws_src.iter_rows(min_row=5, max_row=20, values_only=True):
        if not row or not row[0]:
            continue
        src_by_id[str(row[0])] = row

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "响应表_一一对应"

    headers = [
        "编号",
        "运行单元（原表）",
        "基座模型（原表）",
        "场景与顺序（原表）",
        "种子（原表要求）",
        "数据填报去向（原表）",
        "原表状态",
        "响应状态",
        "已有种子",
        "已有场景",
        "成果物行数/条目",
        "成果物路径",
        "MATLAB归档提示",
        "缺口/备注",
        "核对时间",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    fills = {
        "完整": PatternFill("solid", fgColor="C6EFCE"),
        "部分完整": PatternFill("solid", fgColor="FFEB9C"),
        "进行中": PatternFill("solid", fgColor="BDD7EE"),
        "缺失": PatternFill("solid", fgColor="FFC7CE"),
    }

    ws.append(headers)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    summary_counts = Counter()

    for rec in records:
        rid = rec["id"]
        src = src_by_id.get(rid, [rid, rec["unit"], "", "", "", "", "", "", "", "待采集"])
        info = rec["info"]
        need = rec["need_seeds"]
        expect_n = rec["expect_n"]
        n = int(info.get("n") or 0)
        seeds = list(info.get("seeds") or [])
        scens = list(info.get("scens") or [])
        paths = list(info.get("paths") or [])
        gap = rec.get("gap") or ""
        note = info.get("note") or ""

        # special in-progress
        st = status_for(need, seeds, n, expect_n, issues=gap)
        if rid == "R12" and opt42_n and opt42_n < 8:
            st = "进行中"
            gap = (gap + "；" if gap else "") + f"Ubuntu optimization seed42 已落盘{opt42_n}/8"
        if rid in ("R15", "R16"):
            st = "缺失"
        if rid in ("R13", "R14") and corr42_n and "42" in seeds and len([s for s in seeds if s in ("42", "43", "44")]) >= 2:
            # protocol mismatch note but data present for 43/44 +42
            if not gap:
                st = "部分完整" if "42" in seeds and len(seeds) < 3 else st
                if "协议实现为C0/C1/C2消融（非字面5次独立运行）" not in (gap + note):
                    gap = (gap + "；" if gap else "") + "协议实现为C0/C1/C2消融（非字面5次独立运行）"
        if rid == "R7" and info.get("empty_dev"):
            st = "部分完整"
        if rid == "R6" and info.get("empty_dev"):
            st = "部分完整"
        if rid == "R1" and n >= 12 and set(need).issubset(set(seeds)):
            st = "完整"
        if rid in ("R9", "R10", "R11") and n >= 4:
            # qualitative single-seed satisfied; extra 43/44 is bonus
            st = "完整" if (rid == "R9" and "42" in seeds) or (rid != "R9" and n >= 4) else st
            if rid != "R9" and "42" in seeds:
                st = "完整"

        summary_counts[st] += 1
        path_str = "\n".join(paths) if paths else "—"
        row = [
            rid,
            src[1],
            src[2],
            src[3],
            src[4],
            src[8],
            src[9] if len(src) > 9 else "待采集",
            st,
            ",".join(seeds) if seeds else "—",
            ",".join(scens) if scens else "—",
            n,
            path_str,
            rec.get("matlab_hint") or "—",
            " | ".join(x for x in [gap, note] if x),
            now,
        ]
        ws.append(row)
        r_idx = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(r_idx, col)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r_idx, 8).fill = fills.get(st, PatternFill())

    # widths
    widths = [6, 28, 14, 16, 18, 18, 10, 10, 12, 14, 10, 42, 36, 40, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # summary sheet
    ws2 = wb.create_sheet("汇总与协议说明")
    ws2["A1"] = "实验清单响应表汇总"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2["A2"] = f"生成时间：{now}"
    ws2["A3"] = f"源表：{SRC}"
    ws2["A5"] = "响应状态统计"
    ws2["A5"].font = Font(bold=True)
    ws2.append([])  # won't use append after explicit cells carefully
    r = 6
    for k in ("完整", "部分完整", "进行中", "缺失"):
        ws2.cell(r, 1, k)
        ws2.cell(r, 2, summary_counts.get(k, 0))
        ws2.cell(r, 1).fill = fills.get(k, PatternFill())
        r += 1
    r += 1
    ws2.cell(r, 1, "公共协议对照（源表第1行）").font = Font(bold=True)
    r += 1
    notes = [
        "主种子 42/43/44 = 三次独立重复实验（非开发/测试划分）。",
        "每个主种子内部：DEV = master*100+{1,2,3}（如4201–4203）；TEST = master*100+{4,5,6}（如4204–4206）。",
        "CSV 中 *_dev 列为优化过程快照指标；*_test 列为冻结参数后在测试种子上的复评。",
        "R9 正向顺序数据复用 R1/E0；R10/R11 另有 reverse/shuffled 消融 CSV。",
        "R13/R14 工程实现为 correction 套件 C1/C2（另有 C0 全管线对照），按主种子跑 T4，不是字面“5次独立清空”目录名。",
        "R12 seed42 与 O0 正在 Ubuntu 重跑（逐场景落盘）；43/44 的 O0/O1 已完整。",
        "R15/R16（Qwen/Kimi）尚未开跑。",
        f"Ubuntu 附加：{ubu.get('raw','').splitlines()[:3] if ubu else '无'}",
    ]
    for line in notes:
        ws2.cell(r, 1, line)
        r += 1
    ws2.column_dimensions["A"].width = 100
    ws2.column_dimensions["B"].width = 8

    # also keep original sheet copy for side-by-side
    ws3 = wb.create_sheet("原表副本")
    for row in ws_src.iter_rows(values_only=True):
        ws3.append(list(row))

    for out in (OUT_DESKTOP, OUT_REPO):
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
        print("Wrote", out)

    print("STATUS_COUNTS", dict(summary_counts))
    for rec in records:
        info = rec["info"]
        print(
            rec["id"],
            "n=",
            info.get("n"),
            "seeds=",
            info.get("seeds"),
            "gap=",
            rec.get("gap"),
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

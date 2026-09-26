#!/usr/bin/env python3
"""Patch the current manuscript for protocol wording (no SAH number swap)."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document

WECHAT = Path(
    r"e:\360MoveData\Users\admin\xwechat_files\Jetcai-xiaoxiong_85e0\temp\RWTemp\2026-08\1f059ba6adf5f406df316c093d432dfe\GSMultiAgent A Multi-Agent Loop Framework atop Hermes Agent for Intelligent Design of Guidance Systems(3).docx"
)
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = Path(r"E:\Chj_projects\Papers") / f"GSMultiAgent_protocol_edit_{STAMP}.docx"
BACKUP = Path(r"E:\Chj_projects\Papers") / f"GSMultiAgent_wechat3_backup_{STAMP}.docx"


def set_text(para, new_text: str) -> None:
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(new_text)


REPLACEMENTS = [
    (
        "Across four engagement scenarios, GSMultiAgent attains 9/12 (75%) development-set and 5/12 (41.7%) test-set feasibility within 1–22 MC calls, 10–225× fewer than budget-exhausting baselines and 6–225× fewer than the single agent with 0/12 first-attempt success.",
        "Across four engagement scenarios, GSMultiAgent attains 9/12 (75%) development-set and 5/12 (41.7%) test-set feasibility. Its tabulated development MC-call counts are 1–22 against a shared 225-call ceiling that the standalone optimizers exhaust; these counts are not in identical units (Section 4.1). The single-agent first-success rate is omitted here pending a search-width realignment.",
    ),
    (
        "All methods share a single development-phase budget. One candidate evaluation comprises one 100-trajectory Monte Carlo (MC) batch on each of the three development seeds — i.e., up to 3 MC calls per candidate. Within each outer workflow iteration, Layer 3 performs a search of depth 15 (PPO episodes / candidate evaluations); with the outer-iteration budget fixed at 5 in the main text, a method can therefore evaluate at most 5 × 15 = 75 candidates and consume at most 75 × 3 = 225 development MC calls. Optimization terminates when either the 225-call cap is exhausted or all three development seeds simultaneously satisfy every constraint — the simultaneous three-seed condition prevents overfitting to a single perturbation. This cap of 225 MC calls is applied identically to GSMultiAgent, Direct PPO, Random Search, and CMA-ES, and Hermes Agent operates within the same limit; no method receives a larger simulation allowance than any other.",
        "All methods share a single development-phase ceiling of 225 Monte Carlo (MC) calls, not a requirement that every run consume 225 calls. During search, each MATLAB call uses 10 trajectories per development seed; the 100-trajectory batch is reserved for the held-out test (and for the frozen-parameter final check). One candidate evaluation queries three development seeds (AND logic), so the ceiling is 5 outer iterations × 15 candidates × 3 seeds = 225 calls. For Random Search, CMA-ES, and Direct PPO this ceiling is counted per MATLAB invocation (three real seeds per candidate). For GSMultiAgent the tabulated MC-call column is derived from outer-loop progress (typical recorded values 1 / 7 / 22) and often executes only the first listed seed in MATLAB, filling the other two by bootstrap sampling; those tabulated integers are therefore not interchangeable with the baseline 225. Optimization may stop at the cap or when all three development seeds satisfy every constraint. No method is allowed a larger ceiling than 225.",
    ),
    (
        "First Success: whether all five constraints are satisfied in the first iteration (agent-based methods, including zero-iteration REUSE_HISTORY paths); for the standalone optimizers, which have no first-iteration notion, whether a feasible candidate is found within the budget, which as recorded coincides with development feasibility for all configurations except the single-agent Hermes (0/12 first success vs. 12/12 final development feasibility).",
        "First Success (Table 5, GSMultiAgent and the standalone optimizers): whether the final development-set design is feasible (all five constraints on all three development seeds). For GSMultiAgent this coincides with the recorded development-feasibility bit (9/12), not with “success on the first outer iteration” (only the REUSE / first-pass cells finish in one loop). For Random Search / CMA-ES / Direct PPO it coincides with whether any candidate in the 225-call search was feasible, which in these data equals final development feasibility. The single-agent Hermes column in Table 6 still uses a stricter “all PPO episodes in one outer iteration” rule (0/12 vs. 12/12 final development feasibility) and is being realigned to the same final-development definition; those Table 6 first-success entries are not compared until the realignment is complete.",
    ),
    (
        "Iterations: number of optimization iterations consumed — outer workflow iterations (1–5) for the agent-based methods; candidate evaluations (75) for Random Search and Direct PPO and generations (8) for CMA-ES among the standalone optimizers.",
        "Iterations: a bookkeeping count whose unit differs by method and must not be compared as a raw integer — outer workflow loops (1–5) for GSMultiAgent and Single-Agent; candidate evaluations (up to 75) for Random Search and Direct PPO; CMA-ES generations (8, with λ = 10 and a truncated final generation).",
    ),
    (
        "MC calls: total number of 100-episode Monte Carlo evaluation batches performed during optimization; with three development seeds, one full iteration corresponds to up to 3 MC calls. This is the quantity that directly measures how much of the equal simulation budget (≤ 225) is consumed, and it separates the contribution of memory-based retrieval — which can reduce the number of MC calls to a single call for REUSE_HISTORY paths — from pure optimization efficiency.",
        "MC calls (development): how much of the 225-call ceiling is recorded as consumed. Search uses 10 trajectories per MATLAB call, not 100. For the standalone optimizers the count is the number of real MATLAB invocations (75 candidates × 3 seeds = 225 when the cap is exhausted). For GSMultiAgent the published integers (1–22) are outer-loop-derived and include bootstrap-filled seeds, so a “1 MC call” REUSE cell is not the same computational object as one baseline candidate. Memory reuse can reduce the recorded count; that reduction is an efficiency observation under this accounting, not a claim that MATLAB work is 10–225× smaller in identical units.",
    ),
    (
        "Time/call: Time/call is defined as the per-evaluation simulation time, expressed in minutes per Monte Carlo batch. For GSMultiAgent, it is computed as (wall_clock − orchestration)/MC calls, which isolates the pure simulation cost from the multi-agent orchestration overhead; for the baseline optimizers (CMA-ES, Direct PPO, Random Search, and Single-Agent Hermes), it is recorded directly by the simulation machinery. During the search phase each MC call simulates 10 trajectories per development seed, whereas the 100-trajectory batch is reserved for final verification; per-call time is therefore measured on the 10-trajectory budget actually used in search.",
        "Time/call: minutes per recorded development MC call, using only simulation time (for GSMultiAgent: (wall_clock − orchestration) / MC calls; for Random Search / CMA-ES / Direct PPO: the evaluator’s MATLAB time). This is not wall-clock / MC calls; mixing the two formulas produces the 4.6–15.7 min/call figures sometimes quoted for GSMultiAgent and must not be compared with the 0.03–0.05 min/call baseline values. Search Time/call is measured on the 10-trajectory budget. The Single-Agent Hermes Time/call currently in Table 6 (≈0.40 min) was collected at 100 trajectories per search call and is pending replacement.",
    ),
    (
        "All experiments are run on a single machine (13th Gen Intel Core i7-13700K, 32.0 GB RAM, Ubuntu 24.04.4 LTS). The MATLAB version is R2025a. The foundation model adopted is DeepSeek-V4-Pro; the embedding model is text-embedding-v3.",
        "GSMultiAgent ablations were collected on a workstation (13th Gen Intel Core i7-13700K, 32.0 GB RAM, MATLAB R2025a). The equal-budget Random Search / CMA-ES / Direct PPO / Single-Agent campaign was collected on a separate Ubuntu 24.04 host with MATLAB R2020b. The foundation model is DeepSeek-V4-Pro; the embedding model is text-embedding-v3.",
    ),
    (
        "The main experiment consists of two complementary comparisons, both on the same foundation model, the same automatic MATLAB Engine feedback harness, and the same evaluation protocol: Section 4.2.1 varies the optimization component while the orchestration and memory configurations are fixed; Section 4.2.2 varies the orchestration (single-agent vs. multi-agent cascade) while the optimization component, memory, and automation are fixed.",
        "The main experiment consists of two complementary comparisons on the same foundation model and the same 225-call ceiling: Section 4.2.1 contrasts GSMultiAgent with standalone parameter optimizers (no design-path selection; each scenario × seed cell is cold-started); Section 4.2.2 contrasts the multi-agent cascade with Single-Agent Hermes. The latter is not a single-factor orchestration ablation: GSMultiAgent may select MODIFY_LAW and accumulates the experience base across T1→T4 within a master seed, whereas Single-Agent Hermes only tunes parameters (TUNE_PARAMS) and the standalone optimizers have no experience base.",
    ),
    (
        "Table 5. Comparison of optimization components on T1–T4 (three master seeds, mean ± SD).",
        "Table 5. Comparison of optimization components on T1–T4 (three master seeds, mean ± SD). First success = final development feasibility. MC calls are a shared 225 ceiling, not identical accounting units (Section 4.1). Iterations are not commensurate across methods.",
    ),
    (
        "GSMultiAgent attains its feasibility within 1–22 MC calls — 0.4%–9.8% of the 225-call cap — i.e., roughly 10–225× fewer simulation calls than the budget-exhausting baselines (e.g., 10 vs. 225 MC calls on T1, 3 vs. 225 on T2, 7 vs. 225 on T4; the T2 success relies on nearly zero-iteration memory reuse). Yet its final test-set feasibility (5/12) is the lowest among all configurations: the budget-exhausting baselines show no development-to-test drop (CMA-ES 12/12 → 12/12, Direct PPO 11/12 → 11/12, Random Search 8/12 → 8/12), whereas GSMultiAgent drops from 9/12 on development to 5/12 on test. Early stopping at the first feasible design leaves less margin against the held-out perturbation realizations than a full-budget search; this trade-off between simulation budget and extrapolation robustness characterises the agent’s early-stopping behaviour under the equal-budget protocol.",
        "GSMultiAgent’s tabulated development counts are 1–22 against a 225-call ceiling that the standalone optimizers exhaust. These ratios should not be read as a 10–225× reduction in identical MATLAB work (Section 4.1): T2’s near-zero recorded count uses REUSE from the same master seed’s earlier scenarios, and GSMultiAgent stops once a feasible development design appears, whereas the baselines continue to the ceiling. Post-hoc from the baseline JSON histories, the first feasible candidate (when one exists) costs a mean of 73.9 / 79.2 / 54.0 development MC calls for Random Search / CMA-ES / Direct PPO. GSMultiAgent’s test-set feasibility (5/12) is the lowest among these configurations: the baselines show no development-to-test drop (CMA-ES 12/12 → 12/12, Direct PPO 11/12 → 11/12, Random Search 8/12 → 8/12), whereas GSMultiAgent drops from 9/12 to 5/12. Early stopping versus full-budget search is the intended trade-off under a shared ceiling.",
    ),
    (
        "Table 6 & Figure 12 compare GSMultiAgent with the single-agent Hermes Agent, and Table 7 & Figure 13 reports the test-set constraint satisfaction of all configurations. Both configurations share the 225-call equal-budget cap.",
        "Table 6 & Figure 12 compare GSMultiAgent with Single-Agent Hermes, and Table 7 & Figure 13 reports test-set constraint satisfaction. Both share the 225-call ceiling. Table 6 first-success and Time/call for Hermes still reflect the previous 100-trajectory search and the stricter first-success rule; a 10-trajectory realignment is in progress and those Hermes entries are not used for like-for-like claims until it completes.",
    ),
    (
        "Table 6. Cascade comparison: GSMultiAgent vs. Hermes Agent (single-agent) (three master seeds, mean ± SD).",
        "Table 6. Cascade comparison: GSMultiAgent vs. Hermes Agent (single-agent) (three master seeds, mean ± SD). Hermes first-success and Time/call pending search-width realignment (Section 4.1).",
    ),
    (
        "The single-agent Hermes Agent never satisfies the constraints in the first iteration (first success 0/12), yet it attains development- and test-set feasibility in 12/12 seeds (12/12 [73.5%, 100.0%]; Appendix D) after iterative refinement. GSMultiAgent satisfies all constraints in the first iteration in 9/12 seeds (75.0%; [42.8%, 94.5%]), but its final feasibility (development 9/12 [42.8%, 94.5%], test 5/12 [15.2%, 72.3%]) is lower than that of the single agent. The two configurations differ mainly in first-iteration success and simulation budget: the single agent consumes 138–225 MC calls per scenario (per-scenario means 183–212; approximately 6–225× GSMultiAgent’s budget, a direction consistent across every scenario × seed combination) and 54.9–96.8 min of wall-clock time per seed (on average 1.6–4.7× GSMultiAgent’s per scenario; GSMultiAgent’s per-seed times are more variable, 8.4–104.4 min, and the two ranges overlap).",
        "Single-Agent Hermes attains development- and test-set feasibility in 12/12 seeds after iterative refinement. Its Table 6 first-success rate of 0/12 used a stricter “all episodes in one outer iteration” rule and is not the same event as GSMultiAgent’s 9/12, which equals final development feasibility (only a subset of those nine cells finish in a single outer loop / REUSE). The two systems also differ in design-path support (MODIFY_LAW vs. TUNE_PARAMS only) and in search trajectory width in the published Hermes rows (100 vs. 10). Hermes wall-clock (54.9–96.8 min per seed) and 138–225 recorded MC calls therefore cannot be read as a 6–225× simulation-budget gap in identical units; a 10-trajectory Hermes rerun is in progress.",
    ),
    (
        "The Time/call difference (4.6–15.7 min/call for GSMultiAgent vs. 0.40–0.42 min/call for the single agent) arises because GSMultiAgent embeds LLM diagnosis and adjudication between successive evaluations, whereas the single agent performs batched serial evaluation; the cost structures are therefore not directly comparable.",
        "Wall-clock per recorded MC call for GSMultiAgent (about 4.6–15.7 min if orchestration is left in the numerator) is not the Table 5/6 Time/call statistic, which subtracts orchestration. The published Hermes Time/call of 0.40–0.42 min used 100 search trajectories and is pending replacement. Cost structures remain not directly comparable.",
    ),
    (
        "Budget-wise, GSMultiAgent needs only 1–22 MC calls per scenario (0.4%–9.8% of the 225-call cap), roughly 10–225× fewer than the budget-exhausting optimizers and 6–225× fewer than the single-agent Hermes Agent (138–225 calls).",
        "Budget-wise, GSMultiAgent’s tabulated development counts are 1–22 against a 225-call ceiling; this is a shared cap with different accounting (Section 4.1), not a 10–225× reduction in identical MATLAB work. Single-Agent Hermes recorded 138–225 calls at 100 search trajectories, which is being realigned.",
    ),
    (
        "GSMultiAgent's advantage lies in first success within a minimal budget, not in the feasibility ceiling.",
        "GSMultiAgent’s advantage in these data is reaching a feasible development design without exhausting the 225-call ceiling, not a higher feasibility ceiling.",
    ),
    (
        "the single-agent Hermes Agent reaches 12/12 dev/test after 3.7–4.0 iterations but with 0/12 first success, 138–225 MC calls, and 74–89 min per scenario. The data reveal a three-way trade-off among first success within a minimal budget (9/12 within 1–22 calls), simulation budget, and feasibility ceiling (CMA-ES 12/12):",
        "the single-agent Hermes Agent reaches 12/12 development/test feasibility after iterative refinement; its 0/12 first-success entry is not aligned with GSMultiAgent’s 9/12 final-development definition (Section 4.1). The data reveal a three-way trade-off among stopping early with a feasible development design (9/12, tabulated 1–22 calls), simulation budget, and feasibility ceiling (CMA-ES 12/12):",
    ),
    (
        "Note: First success = all five constraints satisfied in the first optimization iteration. The standalone optimizers involve no agent-level design path and are reported on the pooled T1–T4 record. Under the equal-budget protocol all methods share the 225-call cap; the 0/12 first-success rates of Direct PPO and Random Search under the former five-iteration protocol (≈10–15 MC calls) are attributable to budget limitation rather than algorithmic inability.",
        "Note: For GSMultiAgent and the standalone optimizers in Table D2 / Table 5, first success = final development-set feasibility (all five constraints on all three development seeds). It is not “success on the first outer iteration.” Standalone optimizers have no design path and are pooled over T1–T4. All methods share a 225-call ceiling. The 0/12 first-success rates of Direct PPO and Random Search under the former five-iteration protocol (≈10–15 MC calls) were budget-limited. Single-Agent Hermes first-success in Table 6 is a different, stricter definition pending realignment.",
    ),
]


def patch(doc: Document) -> list[str]:
    hit = []
    for old, new in REPLACEMENTS:
        found = False
        for para in doc.paragraphs:
            if old in para.text:
                set_text(para, para.text.replace(old, new, 1))
                found = True
                hit.append(old[:60])
                break
        if not found:
            hit.append("MISS: " + old[:80])
    return hit


def main() -> int:
    if not WECHAT.is_file():
        raise SystemExit(f"missing {WECHAT}")
    shutil.copy2(WECHAT, BACKUP)
    shutil.copy2(WECHAT, OUT)
    doc = Document(str(OUT))
    report = patch(doc)
    doc.save(str(OUT))
    n_ok = sum(1 for x in report if not x.startswith("MISS"))
    n_miss = sum(1 for x in report if x.startswith("MISS"))
    print(f"saved {OUT}")
    print(f"backup {BACKUP}")
    print(f"replaced {n_ok}  missed {n_miss}")
    for line in report:
        print(line)
    return 0 if n_miss == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
